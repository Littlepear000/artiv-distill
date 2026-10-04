"""Shared logic for report fields / values: slugs, name matching, auto-selection,
Excel import parsing, and parse-result ingestion (used by routes and the Celery hook)."""
import io
import re
import unicodedata
import uuid
from dataclasses import dataclass
from datetime import date, datetime, timezone
from typing import Any, Iterable

from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models.report import Report
from app.models.report_field import ReportField
from app.models.report_field_value import ReportFieldValue

# ---------------------------------------------------------------- pure helpers

_last_ts: datetime | None = None


def next_timestamp() -> datetime:
    """Strictly increasing UTC timestamp so "newest" is well defined within one transaction."""
    global _last_ts
    now = datetime.now(timezone.utc)
    if _last_ts is not None and now <= _last_ts:
        now = _last_ts + datetime.resolution
    _last_ts = now
    return now


def slugify(label: str) -> str:
    text = unicodedata.normalize("NFKD", label or "").encode("ascii", "ignore").decode("ascii").lower()
    text = re.sub(r"[^a-z0-9]+", "_", text).strip("_")
    return text or "field"


_CJK = "一-鿿㐀-䶿"
_NON_ALNUM = re.compile(rf"[^a-z0-9{_CJK}]+")


def normalize_name(name: str) -> str:
    """Lowercase, strip directory + extension, collapse non-alphanumerics (CJK kept) to a space."""
    base = (name or "").replace("\\", "/").rsplit("/", 1)[-1]
    base = re.sub(r"\.[A-Za-z0-9]{1,5}$", "", base)
    return _NON_ALNUM.sub(" ", base.lower()).strip()


def normalize_report_name(name: str) -> str:
    return _NON_ALNUM.sub(" ", (name or "").lower()).strip()


def match_names_to_keys(
    names: dict[str, str], keys: Iterable[str]
) -> tuple[dict[str, str], dict[str, list[str]], list[str], list[str]]:
    """names: {report_id: report_name}. Returns (matched {id: key}, ambiguous {id: [keys]},
    unmatched_ids, unmatched_keys). Exact normalized match first, then unambiguous containment."""
    keys = list(keys)
    by_norm: dict[str, list[str]] = {}
    for k in keys:
        n = normalize_name(k)
        if n:
            by_norm.setdefault(n, []).append(k)

    matched: dict[str, str] = {}
    ambiguous: dict[str, list[str]] = {}
    unmatched: list[str] = []
    for rid, name in names.items():
        n = normalize_report_name(name)
        if not n:
            unmatched.append(rid)
            continue
        exact = by_norm.get(n, [])
        if len(exact) == 1:
            matched[rid] = exact[0]
            continue
        if len(exact) > 1:
            ambiguous[rid] = sorted(exact)
            continue
        partial = sorted(
            k for norm, ks in by_norm.items() if n in norm or norm in n for k in ks
        )
        if len(partial) == 1:
            matched[rid] = partial[0]
        elif len(partial) > 1:
            ambiguous[rid] = partial
        else:
            unmatched.append(rid)
    used = set(matched.values()) | {k for ks in ambiguous.values() for k in ks}
    unmatched_keys = [k for k in keys if k not in used]
    return matched, ambiguous, unmatched, unmatched_keys


def pick_auto_selected(values: list[Any], mode: str):
    """values: objects with .source and .created_at. Returns the one to auto-select, or None."""
    if not values:
        return None
    if mode == "system_first":
        system = [v for v in values if v.source == "system_parse"]
        if system:
            return max(system, key=lambda v: v.created_at)
    return max(values, key=lambda v: v.created_at)


def format_cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == datetime.min.time() else value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


NAME_HEADERS = {"name", "report name", "title"}
URL_HEADERS = {"url", "link"}


@dataclass
class ImportLayout:
    name_col: int | None
    url_col: int | None
    field_cols: dict[int, ReportField]
    unknown_columns: list[str]


def map_headers(headers: list[str], fields: list[ReportField]) -> ImportLayout:
    importable = [f for f in fields if f.allow_import]
    by_label = {f.label.strip().lower(): f for f in importable}
    by_key = {f.key.lower(): f for f in importable}
    name_col = url_col = None
    field_cols: dict[int, ReportField] = {}
    unknown: list[str] = []
    for idx, header in enumerate(headers):
        h = (header or "").strip().lower()
        if not h:
            continue
        if h in NAME_HEADERS and name_col is None:
            name_col = idx
        elif h in URL_HEADERS and url_col is None:
            url_col = idx
        elif h in by_label or h in by_key:
            field_cols[idx] = by_label.get(h) or by_key[h]
        else:
            unknown.append(header.strip())
    return ImportLayout(name_col, url_col, field_cols, unknown)


def read_workbook_rows(data: bytes) -> tuple[list[str], list[tuple[int, list[str]]]]:
    """Returns (headers, [(excel_row_number, cells)]) for the first sheet; fully empty rows dropped."""
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        ws = wb.worksheets[0]
        rows = ws.iter_rows(values_only=True)
        try:
            header_row = next(rows)
        except StopIteration:
            return [], []
        headers = [format_cell(c) for c in header_row]
        out: list[tuple[int, list[str]]] = []
        for i, row in enumerate(rows, start=2):
            cells = [format_cell(c) for c in row]
            if any(cells):
                out.append((i, cells))
        return headers, out
    finally:
        wb.close()


def read_table_rows(data: bytes, filename: str) -> tuple[list[str], list[tuple[int, list[str]]]]:
    """xlsx 或 csv 统一读成 (headers, [(行号, cells)])。csv 依次尝试 utf-8-sig / gb18030。"""
    if not filename.lower().endswith(".csv"):
        return read_workbook_rows(data)
    import csv

    for enc in ("utf-8-sig", "gb18030"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ValueError("Unsupported CSV encoding")
    reader = csv.reader(io.StringIO(text))
    try:
        headers = [h.strip() for h in next(reader)]
    except StopIteration:
        return [], []
    out = [(i, [c.strip() for c in row]) for i, row in enumerate(reader, start=2) if any(c.strip() for c in row)]
    return headers, out


_DATE_RE = re.compile(r"^\d{4}[-/.]\d{1,2}[-/.]\d{1,2}([ T].*)?$|^\d{1,2}[-/.]\d{1,2}[-/.]\d{4}$")


def infer_data_type(header: str, samples: list[str]) -> str:
    """根据列头和样本值猜字段类型：year / percent / number / date / text。"""
    vals = [v.strip() for v in samples if v and v.strip()]
    h = header.strip().lower()
    if not vals:
        return "year" if h in ("year", "年份", "年") else "text"

    def is_num(v: str) -> bool:
        try:
            float(v.replace(",", ""))
            return True
        except ValueError:
            return False

    if all(v.endswith("%") and is_num(v[:-1]) for v in vals):
        return "percent"
    if all(v.isdigit() and len(v) == 4 and 1900 <= int(v) <= 2100 for v in vals) or (
        h in ("year", "年份", "年") and all(is_num(v) for v in vals)
    ):
        return "year"
    if all(is_num(v) for v in vals):
        return "number"
    if all(_DATE_RE.match(v) for v in vals):
        return "date"
    return "text"


def build_template_workbook(fields: list[ReportField]) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Reports"
    ws.append(["name", "url"] + [f.label for f in fields if f.allow_import])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------------ DB helpers


def recompute_selection(db: Session, report_id: uuid.UUID, field: ReportField) -> None:
    """Re-run auto-selection unless the currently selected row is pinned."""
    values = list(
        db.execute(
            select(ReportFieldValue).where(
                ReportFieldValue.report_id == report_id, ReportFieldValue.field_id == field.id
            )
        ).scalars()
    )
    current = next((v for v in values if v.is_selected), None)
    if current is not None and current.is_pinned:
        return
    target = pick_auto_selected(values, field.auto_select)
    if target is current:
        return
    if current is not None:
        current.is_selected = False
        db.flush()
    if target is not None:
        target.is_selected = True
    db.flush()


def touch_report(report: Report) -> None:
    report.updated_at = datetime.now(timezone.utc)


def add_value(
    db: Session,
    *,
    tenant_id: uuid.UUID,
    report: Report,
    field: ReportField,
    value: str,
    source: str,
    note: str | None = None,
    run_id: uuid.UUID | None = None,
    created_by: uuid.UUID | None = None,
    recompute: bool = True,
) -> ReportFieldValue:
    row = ReportFieldValue(
        id=uuid.uuid4(),
        tenant_id=tenant_id,
        project_id=report.project_id,
        report_id=report.id,
        field_id=field.id,
        value=value,
        source=source,
        note=note,
        run_id=run_id,
        created_by=created_by,
        created_at=next_timestamp(),
        is_selected=False,
        is_pinned=False,
    )
    db.add(row)
    db.flush()
    touch_report(report)
    if recompute:
        recompute_selection(db, report.id, field)
    return row


def find_report(db: Session, project_id: uuid.UUID, name: str) -> Report | None:
    name = (name or "").strip()
    if not name:
        return None
    exact = db.execute(select(Report).where(Report.project_id == project_id, Report.name == name)).scalar_one_or_none()
    if exact is not None:
        return exact
    target = normalize_report_name(name)
    if not target:
        return None
    candidates = [
        r for r in db.execute(select(Report).where(Report.project_id == project_id)).scalars()
        if normalize_report_name(r.name) == target
    ]
    return candidates[0] if len(candidates) == 1 else None


def ingest_parse_results(
    db: Session,
    *,
    tenant_id: uuid.UUID,
    project_id: uuid.UUID,
    items: list[dict],
    note: str | None = None,
    run_id: uuid.UUID | None = None,
    created_by: uuid.UUID | None = None,
) -> dict:
    """Append source="system_parse" values. items: [{report_id?|report_name?, fields:{key: value}}]."""
    fields = {
        f.key: f for f in db.execute(select(ReportField).where(ReportField.project_id == project_id)).scalars()
    }
    added = 0
    touched: dict[uuid.UUID, tuple[Report, dict[str, str]]] = {}
    unknown_reports: list[str] = []
    unknown_fields: list[str] = []
    for item in items:
        report = None
        rid = item.get("report_id")
        if rid:
            try:
                report = db.execute(
                    select(Report).where(Report.id == uuid.UUID(str(rid)), Report.project_id == project_id)
                ).scalar_one_or_none()
            except ValueError:
                report = None
        if report is None and item.get("report_name"):
            report = find_report(db, project_id, item["report_name"])
        if report is None:
            label = str(item.get("report_name") or item.get("report_id") or "")
            if label not in unknown_reports:
                unknown_reports.append(label)
            continue
        for key, raw in (item.get("fields") or {}).items():
            field = fields.get(key)
            if field is None or not field.allow_extract:
                if key not in unknown_fields:
                    unknown_fields.append(key)
                continue
            if raw is None or (isinstance(raw, str) and not raw.strip()):
                continue
            add_value(
                db, tenant_id=tenant_id, report=report, field=field, value=format_cell(raw),
                source="system_parse", note=note, run_id=run_id, created_by=created_by,
            )
            added += 1
            touched.setdefault(report.id, (report, {}))[1][key] = format_cell(raw)
    if added:
        # Local import: result_versions imports this module.
        from app.services.result_versions import record_run_version

        record_run_version(
            db, tenant_id=tenant_id, project_id=project_id, fields=fields,
            touched=list(touched.values()), run_id=run_id, created_by=created_by,
        )
    return {"values_added": added, "unknown_reports": unknown_reports, "unknown_fields": unknown_fields}
